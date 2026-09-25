"""Kärkisten vesiosuuskunnan aineisto yhdeksi JSON-tiedostoksi tuontia varten.

    python scripts/karkinen/parse.py <kansio> [--lainat 082026] [--ulos data/private/karkinen/karkinen.json]

Kansiossa on oltava:
  - Ennakkotavoitelista veloituksista YYYYMM.pdf  (kuukauden veloitukset huoneistoittain)
  - Käyttöpaikat*.xlsx                             (kulutuspiste, osoite ja yhteystiedot)
  - Lainaoosuudet*.xlsx                            (lainaosuuksien saldot kuukausittain)
Valinnaiset (asiakastietojen täydennys):
  - Asiakasrekisteri*.xlsx                         (henkilöt: laskutusosoite, sähköposti, puhelin)
  - Sidokset*.xlsx                                 (huoneiston henkilöt ja voimassaolo)
  - Verkkolaskutusosoitteet*.pdf                   (henkilön verkkolaskuosoite ja välittäjä)

Ennakkolista on ensisijainen rekisteri: jokainen huoneisto on yksi laskutettava
kiinteistö ja maksaja. Käyttöpaikka yhdistetään huoneistoon osoitteella, ja jos
osoite ei täsmää, maksajan nimellä. Lainaosuuden saldo otetaan valitun kuukauden
välilehdeltä (oletus 082026 = saldo 31.8.2026), jolloin syyskuun lyhennys ja
korko lasketaan siitä.

Tulostaa vain määriä. Tulos sisältää henkilötietoja, joten se kirjoitetaan
data/private-kansioon, joka ei ole versionhallinnassa.
"""

import collections
import csv
import datetime
import difflib
import glob
import io
import json
import os
import re
import sys

import openpyxl
import pymupdf


def arg(name, default=None):
    return sys.argv[sys.argv.index(name) + 1] if name in sys.argv else default


def num(s):
    return float(s.replace(" ", "").replace(",", "."))


def norm(s):
    return re.sub(r"[^a-z0-9åäö]", "", str(s or "").lower())


def tokens(s):
    return {t for t in re.findall(r"[a-zåäö]{3,}", str(s or "").lower()) if t not in ("oy", "kuolinpesä")}


def find_one(folder, pattern):
    hits = glob.glob(os.path.join(folder, pattern))
    if len(hits) != 1:
        sys.exit(f"Kansiosta pitää löytyä täsmälleen yksi tiedosto {pattern} (löytyi {len(hits)}).")
    return hits[0]


def parse_charges(path):
    """Ennakkolistan huoneistot ja veloitusrivit sarakkeiden x-sijainnin perusteella."""
    doc = pymupdf.open(path)
    units, cur = [], None
    for page in doc:
        rows = collections.defaultdict(list)
        for w in page.get_text("words"):
            rows[round(w[1] / 2)].append(w)
        for key in sorted(rows):
            ws = sorted(rows[key], key=lambda w: w[0])
            txt = [w[4] for w in ws]
            # Huoneistorivi: numero vasemmassa reunassa ja jakson päivämäärä x≈190.
            if ws and 20 <= ws[0][0] <= 30 and re.fullmatch(r"\d+", ws[0][4]) and any(
                abs(w[0] - 190) < 5 and re.match(r"\d\d\.\d\d\.\d{4}", w[4]) for w in ws
            ):
                cur = {
                    "nro": ws[0][4],
                    "osoite": " ".join(w[4] for w in ws if 50 < w[0] < 185),
                    "maksaja": " ".join(w[4] for w in ws if w[0] > 440 and not re.fullmatch(r"\d+", w[4])),
                    "rivit": [],
                    "yhteensa": None,
                }
                units.append(cur)
                continue
            # Veloitusrivi: kaksinumeroinen maksulaji ja jakso x≈139.
            if cur and ws and ws[0][0] < 25 and re.fullmatch(r"\d\d", ws[0][4]) and any(abs(w[0] - 139) < 5 for w in ws):
                qty = [w for w in ws if 320 < w[0] < 360 and re.match(r"[\d,]+$", w[4])]
                price = [w for w in ws if 400 < w[0] < 430]
                vat = [w for w in ws if 485 < w[0] < 515]
                total = [w for w in ws if w[0] > 535]
                cur["rivit"].append({
                    "koodi": ws[0][4],
                    "laji": " ".join(w[4] for w in ws if 30 < w[0] < 135),
                    "maara": num(qty[0][4]) if qty else None,
                    "yks": "m3" if "m3" in txt else None,
                    "hinta": num(price[0][4]) if price else None,
                    "alv": num(vat[0][4]) if vat else 0.0,
                    "summa": num(total[-1][4]) if total else None,
                })
            if cur and "yhteensä" in txt:
                total = [w for w in ws if w[0] > 535]
                if total:
                    cur["yhteensa"] = num(total[-1][4])
    return units


def parse_usage_points(path):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h or "").strip() for h in rows[0]]
    expected = ["Kulutuspiste", "Osoite", "Paikkakunta", "Postinumero", "HUOM.", "HUOM 2.", "Uusi mittari?", "Mittari vaihdettu",
                "Asiakkaan nimi", "Sähköposti", "Puhelinnumero 1", "Puhelinnumero 2"]
    if header[: len(expected)] != expected:
        sys.exit("Käyttöpaikat-taulukon sarakkeet ovat muuttuneet.")
    out = []
    for r in rows[1:]:
        if r[0] is None:
            continue
        s = lambda v: str(v).strip() if v is not None and str(v).strip() else None
        out.append({
            "kulutuspiste": str(r[0]), "osoite": s(r[1]), "paikkakunta": s(r[2]), "postinumero": s(r[3]),
            "huom": " / ".join(x for x in (s(r[4]), s(r[5]), s(r[6])) if x) or None, "mittari_vaihdettu": s(r[7]),
            "nimi": s(r[8]), "email": s(r[9]), "puhelin": s(r[10]), "puhelin2": s(r[11]),
        })
    return out


def parse_loans(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f"Lainaosuuksissa ei ole välilehteä {sheet}.")
    month, year = int(sheet[:2]), int(sheet[2:])
    # Saldo kuukauden viimeisenä päivänä.
    nxt = (year + (month == 12), month % 12 + 1)
    balance_date = (datetime.date(nxt[0], nxt[1], 1) - datetime.timedelta(days=1)).isoformat()
    out = []
    for r in wb[sheet].iter_rows(values_only=True):
        if not r[0] or not isinstance(r[2], (int, float)) or str(r[0]).strip().startswith("Yhteensä"):
            continue
        balance = round(float(r[2]), 2)
        if balance <= 0:
            continue
        note = str(r[8]).strip() if len(r) > 8 and r[8] else None
        final = None
        m = re.search(r"(?:pois|loppuun)\s+(\d{1,2})\.(\d{1,2})\.(\d{4})", note or "", re.I)
        if m:
            final = f"{m.group(3)}-{int(m.group(2)):02d}-01"
        out.append({
            "osoite": str(r[0]).strip(), "osakas": str(r[1] or "").strip(), "saldo": balance, "saldopaiva": balance_date,
            "lyhennys": round(float(r[4]), 2) if isinstance(r[4], (int, float)) else 0.0, "viimeinen_kk": final, "huom": note,
        })
    return out


# Henkilötunnuksia ei käsitellä: vapaatekstistä poistetaan tunnuksen näköiset merkkijonot.
HETU = re.compile(r"\b\d{6}[-+A-FU-Y]\d{3}[0-9A-Y]\b")


def clean_note(text, stats):
    text = (text or "").strip()
    if HETU.search(text):
        stats["henkilötunnus poistettu"] += 1
        text = HETU.sub("[poistettu]", text)
    return text or None


def find_optional(folder, pattern):
    hits = glob.glob(os.path.join(folder, pattern))
    return hits[0] if len(hits) == 1 else None


def read_csv_sheet(path, expected):
    """Taulukko, jonka jokaisella rivillä on yksi pilkuin eroteltu CSV-rivi."""
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    lines = [str(r[0]) for r in ws.iter_rows(values_only=True) if r and r[0]]
    rows = list(csv.reader(io.StringIO("\n".join(lines))))
    if [h.strip() for h in rows[0][: len(expected)]] != expected:
        sys.exit(f"{os.path.basename(path)}: sarakkeet ovat muuttuneet.")
    return rows[1:]


def fi_date(s):
    m = re.match(r"\s*(\d{1,2})\.(\d{1,2})\.(\d{4})", s or "")
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else None


def parse_persons(path, stats):
    out = []
    for r in read_csv_sheet(path, ["Etunimi", "Sukunimi", "Voimassaolevat henkilöroolit", "Lisätiedot", "Sähköposti",
                                   "Matkapuhelin", "Matkapuhelin (norm.)", "Työpuhelin", "Osoite"]):
        if len(r) < 9:
            stats["henkilörivi ohitettu"] += 1
            continue
        # Pilkku osoitteessa ("Katu 1, 41800 Korpilahti") voi jakaa rivin ilman lainausmerkkejä.
        address = ",".join(r[8:]).strip()
        m = re.match(r"^(.*?),?\s*(\d{5})\s+(.+)$", address)
        phone = r[6].strip() if r[6].strip().startswith("+") else r[5].strip()
        out.append({
            "etunimi": r[0].strip(), "sukunimi": r[1].strip(), "email": r[4].strip() or None, "puhelin": phone or None,
            "katu": m.group(1).strip() if m else (address or None), "postinumero": m.group(2) if m else None,
            "toimipaikka": m.group(3).strip() if m else None, "lisatiedot": clean_note(r[3], stats),
        })
    return out


def parse_bonds(path, stats):
    out = []
    for r in read_csv_sheet(path, ["Huoneisto", "Henkilö", "Sidosnumero", "Lisähenkilö", "Sidostyyppi", "Voimassa", "Merkitty", "Asuu", "Lisätiedot"]):
        if len(r) < 6:
            stats["sidosrivi ohitettu"] += 1
            continue
        valid = r[5].split("-")
        out.append({
            "huoneisto": r[0].strip(), "henkilo": r[1].strip(), "sidosnumero": r[2].strip(), "lisahenkilo": r[3].strip() or None,
            "alku": fi_date(valid[0]), "loppu": fi_date(valid[1]) if len(valid) > 1 else None,
            "lisatiedot": clean_note(",".join(r[8:]), stats) if len(r) > 8 else None,
        })
    return out


def parse_einvoices(path, stats):
    """Verkkolaskuosoitteet: nimi vasemmalla (voi olla omalla rivillään), välittäjä x≈481, osoite x≈593."""
    out, pending = [], []
    for page in pymupdf.open(path):
        rows = collections.defaultdict(list)
        for w in page.get_text("words"):
            rows[round(w[1] / 3)].append(w)
        for key in sorted(rows):
            ws = sorted(rows[key], key=lambda w: w[0])
            words = [w[4] for w in ws]
            if any(x in words for x in ("Tulostettu", "Sivu", "Välittäjätunnus")):
                continue
            name = [w[4] for w in ws if w[0] < 225]
            operator = [w[4] for w in ws if 475 < w[0] < 585]
            address = [w[4] for w in ws if 585 < w[0] < 700]
            direct = [w[4] for w in ws if w[0] >= 700]
            if operator and address:
                if not name and not pending:
                    stats["verkkolasku ilman nimeä"] += 1
                    continue
                if name and pending:
                    stats["verkkolasku, edellinen nimirivi ohitettu"] += 1
                out.append({"nimi": " ".join(name or pending), "valittaja": operator[0], "osoite": address[0],
                            "suoramaksu": bool(direct) and direct[0] == "Kyllä"})
                pending = []
            elif name:
                pending = name
    return out


def person_tokens(p):
    # Sukunimi ja ensimmäinen etunimi.
    first = re.findall(r"[a-zåäö]{3,}", (p["etunimi"] or "").lower())[:1]
    return tokens(p["sukunimi"]) | set(first)


def overlap(a, b):
    return bool(a and b and len(a & b) >= min(2, len(a), len(b)))


def match_registry(folder, units, period_end, stats):
    """Henkilö-, sidos- ja verkkolaskutiedot huoneistoille (kenttä "rekisteri")."""
    persons_path = find_optional(folder, "Asiakasrekisteri*.xlsx")
    bonds_path = find_optional(folder, "Sidokset*.xlsx")
    einv_path = find_optional(folder, "Verkkolaskutusosoitteet*.pdf")
    persons = parse_persons(persons_path, stats) if persons_path else []
    bonds = parse_bonds(bonds_path, stats) if bonds_path else []
    einvoices = parse_einvoices(einv_path, stats) if einv_path else []
    stats["henkilöitä"], stats["sidoksia"], stats["verkkolaskuosoitteita"] = len(persons), len(bonds), len(einvoices)

    # Sidos huoneistolle osoitteella; samassa osoitteessa useampi huoneisto (39, 160) → henkilön nimellä.
    by_addr = collections.defaultdict(list)
    for u in units:
        u["rekisteri"] = {"sidokset": [], "maksajan_sidos": None, "henkilo": None, "verkkolasku": None}
        by_addr[norm(u["osoite"])].append(u)
    for b in bonds:
        c = by_addr.get(norm(b["huoneisto"]), [])
        if not c:
            c = [u for u in units if difflib.SequenceMatcher(None, norm(b["huoneisto"]), norm(u["osoite"])).ratio() > 0.9]
        if len(c) > 1:
            named = [u for u in c if overlap(tokens(b["henkilo"]), tokens(u["maksaja"]))]
            c = named if len(named) == 1 else c
        for u in c:
            u["rekisteri"]["sidokset"].append(b)
        stats["sidos yhdistetty" if len(c) == 1 else "sidos useaan huoneistoon" if c else "sidos ilman huoneistoa"] += 1

    for u in units:
        reg = u["rekisteri"]
        current = [b for b in reg["sidokset"] if b["loppu"] is None or b["loppu"] >= period_end]
        payer = [b for b in current if overlap(tokens(b["henkilo"]), tokens(u["maksaja"]))]
        reg["maksajan_sidos"] = payer[0] if len(payer) == 1 else (current[0] if len(current) == 1 else None)
        if reg["maksajan_sidos"]:
            stats["maksajan sidos"] += 1
        # Henkilö: sidoksen nimellä, muuten maksajan nimellä.
        sources = ([reg["maksajan_sidos"]["henkilo"]] if reg["maksajan_sidos"] else []) + [u["maksaja"]]
        for source in sources:
            t = tokens(source)
            c = [p for p in persons if len(person_tokens(p)) >= 2 and person_tokens(p) <= t]
            if len(c) == 1:
                reg["henkilo"] = c[0]
                stats["henkilö yhdistetty"] += 1
                break

    for e in einvoices:
        t = tokens(e["nimi"])
        free = [u for u in units if u["rekisteri"]["verkkolasku"] is None]
        c = [u for u in free if overlap(t, tokens(u["maksaja"]))]
        if len(c) != 1:
            c = [u for u in free if u["rekisteri"]["henkilo"] and person_tokens(u["rekisteri"]["henkilo"]) <= t]
        if len(c) == 1:
            c[0]["rekisteri"]["verkkolasku"] = e
            stats["verkkolasku yhdistetty"] += 1
        else:
            stats["verkkolasku ei yhdistetty"] += 1


def main():
    if len(sys.argv) < 2 or sys.argv[1].startswith("--"):
        sys.exit(__doc__)
    folder = sys.argv[1]
    charges_pdf = find_one(folder, "Ennakkotavoitelista*.pdf")
    period = re.search(r"(\d{6})", os.path.basename(charges_pdf)).group(1)
    units = parse_charges(charges_pdf)
    points = parse_usage_points(find_one(folder, "Käyttöpaikat*.xlsx"))

    # Sama huoneistonumero voi esiintyä kahdesti eri maksajalla (esim. 39 ja 160):
    # jälkimmäinen saa tunnuksen "39-2", jotta kumpikin on oma laskutuskohteensa.
    seen = collections.Counter()
    for u in units:
        seen[u["nro"]] += 1
        u["tunnus"] = u["nro"] if seen[u["nro"]] == 1 else f"{u['nro']}-{seen[u['nro']]}"
    loans = parse_loans(find_one(folder, "Lainaoosuudet*.xlsx"), arg("--lainat", "082026"))

    # Käyttöpaikka huoneistolle: ensin yksiselitteinen osoite, sitten maksajan nimi.
    by_addr = collections.defaultdict(list)
    for p in points:
        by_addr[norm(p["osoite"])].append(p)
    used, stats = set(), collections.Counter()
    for u in units:
        c = [p for p in by_addr.get(norm(u["osoite"]), []) if p["kulutuspiste"] not in used]
        if len(c) == 1:
            u["kayttopaikka"] = c[0]
            used.add(c[0]["kulutuspiste"])
            stats["osoitteella"] += 1
    for u in units:
        if "kayttopaikka" in u:
            continue
        t = tokens(u["maksaja"])
        cands = [p for p in points if p["kulutuspiste"] not in used and t and len(t & tokens(p["nimi"])) >= min(2, len(t))]
        same_road = [p for p in cands if norm(u["osoite"])[:6] == norm(p["osoite"])[:6]]
        pick = same_road if len(same_road) == 1 else cands
        if len(pick) == 1:
            u["kayttopaikka"] = pick[0]
            used.add(pick[0]["kulutuspiste"])
            stats["nimellä"] += 1
        else:
            u["kayttopaikka"] = None
            stats["ei yhdistetty"] += 1

    # Lainaosuus huoneistolle osoitteella, muuten osakkaan nimellä.
    unit_by_addr = collections.defaultdict(list)
    for u in units:
        u["laina"] = None
        unit_by_addr[norm(u["osoite"])].append(u)
    loan_miss = 0
    for loan in loans:
        c = [u for u in unit_by_addr.get(norm(loan["osoite"]), []) if u["laina"] is None]
        if len(c) != 1:
            t = tokens(loan["osakas"])
            c = [u for u in units if u["laina"] is None and t and len(t & tokens(u["maksaja"])) >= min(2, len(t))]
        if len(c) != 1:
            # Kirjoitusasu vaihtelee (esim. "Kärkistenlaiturintie" / "Kärkisten laiturintie").
            c = [u for u in units if u["laina"] is None and difflib.SequenceMatcher(None, norm(loan["osoite"]), norm(u["osoite"])).ratio() > 0.9]
        if len(c) == 1:
            c[0]["laina"] = loan
        else:
            loan_miss += 1

    reg_stats = collections.Counter()
    year, month = int(period[:4]), int(period[4:])
    period_end = (datetime.date(year + (month == 12), month % 12 + 1, 1) - datetime.timedelta(days=1)).isoformat()
    match_registry(folder, units, period_end, reg_stats)

    out = arg("--ulos", "data/private/karkinen/karkinen.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        json.dump({"jakso": period, "huoneistot": units, "kayttopaikat_ilman_huoneistoa": [p for p in points if p["kulutuspiste"] not in used]},
                  f, ensure_ascii=False, indent=1)
    total = round(sum(r["summa"] or 0 for u in units for r in u["rivit"]), 2)
    print(f"Jakso {period}: huoneistoja {len(units)}, veloitusrivejä {sum(len(u['rivit']) for u in units)}, yhteensä {total} €.")
    print(f"Käyttöpaikkoja {len(points)}: yhdistetty {stats['osoitteella']} osoitteella ja {stats['nimellä']} nimellä, "
          f"huoneistoja ilman käyttöpaikkaa {stats['ei yhdistetty']}.")
    dups = sorted({u["nro"] for u in units if u["tunnus"] != u["nro"]}, key=int)
    if dups:
        print(f"Sama huoneistonumero useammalla maksajalla: {', '.join(dups)}.")
    print(f"Lainaosuuksia avoinna {len(loans)}, yhdistetty {len(loans) - loan_miss}, ei yhdistetty {loan_miss}.")
    if reg_stats:
        print("Rekisteri: " + ", ".join(f"{k} {v}" for k, v in reg_stats.items()) + ".")
    print(f"Kirjoitettu {out}")


if __name__ == "__main__":
    main()
